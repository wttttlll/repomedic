
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


NAVY = RGBColor(6, 30, 85)
BLUE = RGBColor(36, 61, 103)
GOLD = RGBColor(255, 188, 0)
WHITE = RGBColor(255, 255, 255)


def add_text(slide, x, y, w, h, text, size, color,
             bold=False):
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    frame = shape.text_frame
    frame.word_wrap = True
    p = frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return shape


def add_rect(slide, x, y, w, h, color):
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def create_outline_ppt(output_path="output.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    # Top title bar
    add_rect(slide, 0, 0, 13.333, 0.75, NAVY)
    add_text(
        slide, 5.3, 0.12, 3, 0.5,
        "Outline", 26, WHITE, True
    )

    sections = [
        "Research Background",
        "SBDO-Based Hull Form Optimization",
        "GeneHull Intelligent Design Platform",
        "AI Agents for Autonomous Hull Design",
        "Summary and Outlook",
    ]

    for i, title in enumerate(sections):
        y = 1.3 + i * 1.05

        add_rect(slide, 2.3, y, 0.8, 0.78, GOLD)
        add_rect(slide, 3.1, y, 8.2, 0.78, BLUE)

        add_text(
            slide, 2.52, y + 0.12, 0.5, 0.5,
            str(i + 1), 24, NAVY, True
        )

        add_text(
            slide, 3.22, y + 0.16, 7.9, 0.45,
            title, 19, WHITE, True
        )

    output = Path(output_path)
    output.parent.mkdir(
        parents=True, exist_ok=True
    )

    prs.save(str(output))
    return str(output)


if __name__ == "__main__":
    create_outline_ppt()