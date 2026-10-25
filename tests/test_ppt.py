
from pathlib import Path

from pptx import Presentation
from integrations.ppt.generator import generate_presentation


def test_generate_ppt():
    """测试 PPT 文件能否正常生成。"""

    base_dir = Path(__file__).resolve().parents[1]

    result = generate_presentation(
        base_dir,
        filename="test_ai.pptx",
        palette="ocean",
        preview=False,
        slides=[
            {
                "layout": "cover",
                "title": "人工智能发展报告",
                "subtitle": "2026 年技术趋势",
            },
            {
                "layout": "bullets",
                "title": "人工智能核心趋势",
                "bullets": [
                    "大语言模型",
                    "多模态人工智能",
                    "智能体技术",
                ],
            },
        ],
    )

    # 检查实际生成的文件
    output_file = base_dir / "output" / "ppt" / "test_ai.pptx"

    assert output_file.exists(), "PPT 文件未生成"

    # 检查 PowerPoint 文件能否打开
    prs = Presentation(str(output_file))

    assert len(prs.slides) == 2, "PPT 页数不正确"

    print("\nPPT 生成成功！")
    print("文件位置：", output_file)
    print("生成结果：", result)


def test_generate_multiple_layouts():
    """测试多种幻灯片版式。"""

    base_dir = Path(__file__).resolve().parents[1]

    generate_presentation(
        base_dir,
        filename="test_layouts.pptx",
        preview=False,
        slides=[
            {
                "layout": "cover",
                "title": "PPT 版式测试",
            },
            {
                "layout": "section",
                "title": "第一章节",
            },
            {
                "layout": "two_column",
                "title": "双栏对比",
                "left": ["方案 A", "成本较低"],
                "right": ["方案 B", "性能更好"],
            },
            {
                "layout": "cards",
                "title": "应用场景",
                "cards": [
                    {
                        "title": "教育",
                        "body": "智能教学",
                    },
                    {
                        "title": "医疗",
                        "body": "辅助诊断",
                    },
                ],
            },
        ],
    )

    output_file = base_dir / "output" / "ppt" / "test_layouts.pptx"

    assert output_file.exists()

    prs = Presentation(str(output_file))

    assert len(prs.slides) == 4