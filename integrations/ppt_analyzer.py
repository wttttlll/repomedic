import json
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


def extract_color(color):
    try:
        if color is not None and color.type is not None:
            if color.rgb is not None:
                return str(color.rgb)
    except (AttributeError, TypeError, ValueError):
        pass
    return None


def analyze_ppt(reference_path, output_path):
    reference_path = Path(reference_path)
    output_path = Path(output_path)

    prs = Presentation(str(reference_path))

    result = {
        "slide_width": prs.slide_width / Inches(1),
        "slide_height": prs.slide_height / Inches(1),
        "slides": [],
        "font_statistics": {},
    }

    fonts = Counter()
    sizes = Counter()

    for slide_index, slide in enumerate(prs.slides):
        slide_info = {
            "page": slide_index + 1,
            "text_boxes": [],
            "tables": [],
            "images": [],
        }

        for shape in slide.shapes:
            position = {
                "x": round(shape.left / Inches(1), 4),
                "y": round(shape.top / Inches(1), 4),
                "w": round(shape.width / Inches(1), 4),
                "h": round(shape.height / Inches(1), 4),
            }

            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if not run.text.strip():
                            continue

                        font = run.font

                        info = {
                            "text": run.text,
                            "font_name": font.name,
                            "font_size": (
                                font.size.pt
                                if font.size is not None
                                else None
                            ),
                            "bold": font.bold,
                            "color": extract_color(font.color),
                            "position": position,
                        }

                        slide_info["text_boxes"].append(info)

                        if font.name:
                            fonts[font.name] += 1

                        if font.size:
                            sizes[font.size.pt] += 1

            if shape.has_table:
                table_info = {
                    "position": position,
                    "rows": len(shape.table.rows),
                    "columns": len(shape.table.columns),
                    "column_widths": [
                        round(c.width / Inches(1), 4)
                        for c in shape.table.columns
                    ],
                    "row_heights": [
                        round(r.height / Inches(1), 4)
                        for r in shape.table.rows
                    ],
                }
                slide_info["tables"].append(table_info)

            if shape.shape_type == 13:
                slide_info["images"].append(position)

        result["slides"].append(slide_info)

    result["font_statistics"] = {
        "fonts": dict(fonts.most_common()),
        "sizes": {
            str(k): v for k, v in sizes.most_common()
        },
    }

    output_path.parent.mkdir(parents=True, exist_ok=True)

    output_path.write_text(
        json.dumps(result, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    return result


if __name__ == "__main__":
    base = Path(__file__).resolve().parents[1]

    analyze_ppt(
        base / "templates/ppt/reference.pptx",
        base / "templates/ppt/style_rules.json",
    )

    print("参考 PPT 样式提取完成")