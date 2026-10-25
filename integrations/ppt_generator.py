"""Structured PPTX generation for RepoMedic; Python 3.10+, python-pptx.

Template files are selected by basename from BASE/templates/ppt only.
Existing template slides are removed by default, while masters, layouts and theme remain.
"""
from __future__ import annotations
import re
import shutil
import subprocess
import tempfile
import uuid
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE

PALETTES = {
    'ocean': ('102A43', '147D92', 'F7FAFC', '243B53', 'F4B860'),
    'forest': ('153D31', '278367', 'F6FAF6', '234238', 'F0B95B'),
    'sunset': ('512B44', 'DD6B55', 'FFF9F5', '463644', 'F1B24A'),
    'minimal': ('182230', '456B8C', 'FFFFFF', '344054', 'E7A94E'),
    "academic_navy": ("061E55", "243D67", "FFFFFF", "243D67", "FFBC00"),
}
LAYOUTS = {"cover","outline","section","bullets","two_column","cards","quote","chart","timeline"}
MAX_SLIDES = 30
MAX_TEXT = 3000


def _rgb(value):
    if not re.fullmatch(r'[0-9A-Fa-f]{6}', value):
        raise ValueError('invalid hex color')
    return RGBColor.from_string(value)


def _safe_filename(name, suffix):
    if not isinstance(name, str) or not re.fullmatch(r'[\w\- ]{1,80}' + re.escape(suffix), name, re.UNICODE):
        raise ValueError(f'文件名必须是不含路径的 {suffix} 文件名')
    return name


def _list_of_text(x, limit=10):
    if not isinstance(x, list) or len(x) > limit or any(not isinstance(v,str) or len(v)>300 for v in x):
        raise ValueError('文本条目格式不正确或超出限制')
    return x


def _textbox(slide, x,y,w,h, text, size, color, bold=False, align=None):
    if not isinstance(text,str) or len(text)>MAX_TEXT:
        raise ValueError('文本过长')
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.02)
    tf.margin_top = tf.margin_bottom = Inches(.02)
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = 'Microsoft YaHei'
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = _rgb(color)
    if align is not None:
        p.alignment = align
    return box


def _rect(slide,x,y,w,h,fill):
    shp=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb=_rgb(fill)
    shp.line.fill.background()
    return shp


def _blank_layout(prs):
    for l in prs.slide_layouts:
        if l.name.lower() in ('blank','空白'):
            return l
    # If no blank layout exists, use the layout with the fewest placeholders.
    return min(prs.slide_layouts,key=lambda l:len(l.placeholders))


def _remove_all_slides(prs):
    slides=prs.slides
    for sid in list(slides._sldIdLst):
        slides.part.drop_rel(sid.rId)
        slides._sldIdLst.remove(sid)


def _background(slide, color, use_template):
    if use_template:
        return
    bg=slide.background.fill; bg.solid(); bg.fore_color.rgb=_rgb(color)


def _bullets(slide, bullets, x,y,w,h, color, size=21):
    bullets=_list_of_text(bullets,8)
    if not bullets:
        return
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame; tf.clear(); tf.word_wrap=True
    for i,text in enumerate(bullets):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text='•  '+text
        p.font.name='Microsoft YaHei'; p.font.size=Pt(size)
        p.font.color.rgb=_rgb(color)
        p.space_after=Pt(13)


def _chart(slide, spec, x,y,w,h, primary, ink):
    if not isinstance(spec,dict): raise ValueError('chart 必须是对象')
    kind=spec.get('type','column')
    types={'column':XL_CHART_TYPE.COLUMN_CLUSTERED,
           'bar':XL_CHART_TYPE.BAR_CLUSTERED,
           'line':XL_CHART_TYPE.LINE_MARKERS,
           'pie':XL_CHART_TYPE.PIE}
    if kind not in types: raise ValueError('不支持的图表类型')
    categories=_list_of_text(spec.get('categories',[]),12)
    series=spec.get('series',[])
    if not categories or not isinstance(series,list) or not 1<=len(series)<=5:
        raise ValueError('图表缺少 categories/series')
    data=CategoryChartData(); data.categories=categories
    for item in series:
        if not isinstance(item,dict) or not isinstance(item.get('name'),str):
            raise ValueError('图表 series 格式错误')
        vals=item.get('values')
        if not isinstance(vals,list) or len(vals)!=len(categories) or any(type(v) not in (int,float) or not -1e9<=v<=1e9 for v in vals):
            raise ValueError('图表数值无效')
        data.add_series(item['name'][:80], vals)
    if kind=='pie' and len(series)!=1: raise ValueError('饼图仅支持一个数据序列')
    chart=slide.shapes.add_chart(types[kind],Inches(x),Inches(y),Inches(w),Inches(h),data).chart
    chart.has_legend=(kind!='pie' and len(series)>1)
    if chart.has_legend: chart.legend.position=XL_LEGEND_POSITION.BOTTOM
    try:
        chart.plots[0].series[0].format.fill.solid()
        chart.plots[0].series[0].format.fill.fore_color.rgb=_rgb(primary)
    except (AttributeError,ValueError):
        pass
    return chart


def _slide(prs, item, idx, total, palette, use_template, custom_layouts):
    if not isinstance(item,dict): raise ValueError('slide 必须是对象')
    kind=item.get('layout','bullets')
    if kind not in LAYOUTS: raise ValueError(f'未知布局 {kind}')
    title=item.get('title','')
    if not isinstance(title,str) or len(title)>120: raise ValueError('标题过长')
    layout_name=item.get('template_layout')
    if layout_name is not None and layout_name not in custom_layouts:
        raise ValueError(f'模板中无此版式: {layout_name}')
    layout=custom_layouts[layout_name] if layout_name else _blank_layout(prs)
    s=prs.slides.add_slide(layout)
    # Remove only generated placeholders, retaining layout/master artwork.
    for ph in list(s.placeholders):
        element=ph._element
        element.getparent().remove(element)
    dark,primary,bg,ink,accent=palette
    _background(s,bg,use_template)
    W=prs.slide_width/Inches(1); H=prs.slide_height/Inches(1)
    sx=W/13.333; sy=H/7.5
    def tx(x,y,w,h,text,size,color,bold=False,align=None):
        return _textbox(s,x*sx,y*sy,w*sx,h*sy,text,size,color,bold,align)
    def rect(x,y,w,h,color): return _rect(s,x*sx,y*sy,w*sx,h*sy,color)
    def bl(items,x,y,w,h,size=21): return _bullets(s,items,x*sx,y*sy,w*sx,h*sy,ink,size)
    subtitle=item.get('subtitle','')
    if not isinstance(subtitle,str) or len(subtitle)>250: raise ValueError('副标题过长')
    if kind=='cover':
        if not use_template: rect(0,0,13.333,7.5,dark)
        rect(.8,1.2,.12,4.4,accent)
        tx(1.2,1.8,10.7,1.9,title,38,(dark if use_template else 'FFFFFF'),True)
        tx(1.2,4.35,10,1.0,subtitle,21,(ink if use_template else 'D9E6ED'))
    elif kind == "outline":
        # Navy title bar
        rect(0, 0, 13.333, 0.75, dark)
        tx(4.7, 0.12, 3.9, 0.48,"Outline", 26, "FFFFFF",True, PP_ALIGN.CENTER)
        sections = _list_of_text(item.get("bullets", []), 5)
        if not sections:
            raise ValueError("outline requires 1-5 sections")
        for i, section in enumerate(sections):
            y = 1.3 + i * 1.05
            rect(2.3, y, 0.8, 0.78, accent)
            rect(3.1, y, 8.2, 0.78, primary)
            tx(2.3, y + 0.12, 0.8, 0.52,str(i + 1), 24, dark,True, PP_ALIGN.CENTER)
            tx(3.22, y + 0.14, 7.9, 0.5,section, 18, "FFFFFF", True)
    elif kind=='section':
        if not use_template: rect(0,0,13.333,7.5,dark)
        tx(1.15,2.4,10.8,1.3,title,35,(dark if use_template else 'FFFFFF'),True)
        tx(1.15,4.0,10.8,1.2,subtitle,19,(ink if use_template else 'D9E6ED'))
    else:
        if palette == PALETTES["academic_navy"]:
            rect(0, 0, 13.333, 0.76, dark)
            tx(0.65, 0.10, 12.0, 0.51,title, 25, "FFFFFF", True)
        else:
            rect(.75, .73, .12, .64, primary)
            tx(1.03, .62, 11.6, .86,title, 28, dark, True)
            rect(.8, 1.58, 11.72, .018,accent)
        if kind=='bullets':
            bl(item.get('bullets',[]),1.05,2.0,11,4.5,22)
        elif kind=='two_column':
            left=_list_of_text(item.get('left',[]),7); right=_list_of_text(item.get('right',[]),7)
            rect(.95,2,5.5,.09,primary); rect(6.85,2,5.5,.09,accent)
            bl(left,1.1,2.35,5.1,4.3,18); bl(right,7,2.35,5.1,4.3,18)
        elif kind=='cards':
            cards=item.get('cards',[])
            if not isinstance(cards,list) or not 1<=len(cards)<=4: raise ValueError('cards 数量需为 1～4')
            width=11.5/len(cards)
            for j,c in enumerate(cards):
                if not isinstance(c,dict): raise ValueError('card 必须是对象')
                ct=str(c.get('title',''))[:70]; body=str(c.get('body',''))[:200]
                x=.9+j*width; rect(x,2.1,width-.2,3.9,'E8F1F4' if not use_template else 'F1F3F5')
                rect(x,2.1,width-.2,.08,primary if j%2==0 else accent)
                tx(x+.18,2.45,width-.58,.6,ct,19,dark,True)
                tx(x+.18,3.3,width-.58,2.3,body,15,ink)
        elif kind=='quote':
            tx(1.15,2,10.9,3.3,'“ '+str(item.get('quote',''))[:350]+' ”',29,dark,True,PP_ALIGN.CENTER)
            tx(1.2,5.6,10.9,.6,subtitle,17,primary,False,PP_ALIGN.RIGHT)
        elif kind=='chart':
            _chart(s,item.get('chart',{}),1.0*sx,1.85*sy,11.4*sx,4.8*sy,primary,ink)
        elif kind=='timeline':
            entries=item.get('events',[])
            if not isinstance(entries,list) or not 2<=len(entries)<=5: raise ValueError('events 需 2～5 项')
            rect(1.3,3.45,10.65,.05,primary)
            for j,e in enumerate(entries):
                if not isinstance(e,dict): raise ValueError('event 必须是对象')
                x=1.25+j*(10.45/(len(entries)-1))
                rect(x,3.3,.25,.35,accent)
                tx(x-.35,2.22,2.25,.7,str(e.get('label',''))[:35],16,dark,True)
                tx(x-.35,3.85,2.25,1.8,str(e.get('text',''))[:95],13,ink)
        tx(11.75,7.03,.7,.3,f'{idx}/{total}',10,ink,False,PP_ALIGN.RIGHT)
    return s


def inspect_template(base: Path, template: str | None = None):
    template_dir=(base/'templates'/'ppt').resolve()
    if template is None:
        if not template_dir.is_dir(): return {'templates':[]}
        return {'templates':sorted(p.name for p in template_dir.glob('*.pptx') if p.is_file() and not p.is_symlink())}
    name=_safe_filename(template,'.pptx')
    path=template_dir/name
    if not path.is_file() or path.is_symlink() or path.stat().st_size>30*1024*1024:
        raise ValueError('模板不存在、大小超过 30 MB 或不安全')
    prs=Presentation(str(path))
    return {'name':name,'slide_count':len(prs.slides),
            'slide_size_inches':[round(prs.slide_width/Inches(1),2),round(prs.slide_height/Inches(1),2)],
            'layouts':list(dict.fromkeys(l.name for l in prs.slide_layouts))}


def generate_presentation(
    base: Path,
    *,
    filename: str,
    slides: list,
    palette: str = "ocean",
    template: str | None = None,
    template_mode: str = "replace",
    preview: bool = True
):
    filename=_safe_filename(filename,'.pptx')
    if palette not in PALETTES: raise ValueError('未知配色')
    if template_mode not in ('replace','append',"reference_fill",): raise ValueError("不支持的模板模式")
    if not isinstance(slides,list) or not 1<=len(slides)<=MAX_SLIDES:
        raise ValueError('slides 必须是 1～30 项数组')
    if len(str(slides).encode('utf-8'))>180_000: raise ValueError('内容超过 180KB')
    template_info=inspect_template(base,template) if template else None
    template_path=base/'templates'/'ppt'/template if template else None
    prs=Presentation(str(template_path)) if template else Presentation()
    if template and template_mode=='replace': _remove_all_slides(prs)
    if template_mode == "reference_fill":
        if template != "base.pptx":
            raise ValueError("reference_fill 必须使用 base.pptx")
        from integrations.ppt_template_filler import (generate_from_reference,)

        return generate_from_reference(
            base=base,
            filename=filename,
            slides=slides,
            template_name=template,
            reference_name="reference.pptx",
            style_name="style_rules.json",
            preview=preview,
        )
    if not template:
        prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    layout_map={l.name:l for l in prs.slide_layouts}
    warnings=[]
    if template and not any(l.name.lower() in ('blank','空白') for l in prs.slide_layouts):
        warnings.append('模板不含 Blank/空白版式：继承的背景元素可能与内容重叠')
    for idx,item in enumerate(slides,1):
        if len(str(item.get('bullets',[])))>1200: warnings.append(f'第 {idx} 页条目较长，建议预览检查')
        _slide(prs,item,idx,len(slides),PALETTES[palette],bool(template),layout_map)
    out=base/'output'/'ppt'; out.mkdir(parents=True,exist_ok=True)
    target=out/filename
    if target.is_symlink(): raise ValueError('输出路径不能为软链接')
    tmp=out/f'.{uuid.uuid4().hex}.pptx'
    try:
        prs.save(str(tmp)); tmp.replace(target)
    finally: tmp.unlink(missing_ok=True)
    result={'success':True,'pptx':str(target),'slides_created':len(slides),
            'total_slides':len(prs.slides),'template':template,'warnings':warnings,'preview':[]}
    if preview:
        soffice=shutil.which('libreoffice') or shutil.which('soffice')
        pdftoppm=shutil.which('pdftoppm')
        if not soffice or not pdftoppm:
            warnings.append('未安装 LibreOffice 或 pdftoppm，跳过预览渲染')
        else:
            preview_dir=out/(target.stem+'_preview')
            preview_dir.mkdir(parents=True,exist_ok=True)
            try:
                with tempfile.TemporaryDirectory() as profile:
                    cmd=[soffice,'-env:UserInstallation=file://'+profile,'--headless',
                         '--convert-to','pdf','--outdir',str(preview_dir),str(target)]
                    p=subprocess.run(cmd,capture_output=True,text=True,timeout=90)
                pdf=preview_dir/(target.stem+'.pdf')
                if p.returncode!=0 or not pdf.exists(): raise RuntimeError(p.stderr[-500:])
                cmd=[pdftoppm,'-f','1','-l',str(min(len(prs.slides),30)),'-jpeg',
                     '-scale-to','1280',str(pdf),str(preview_dir/'slide')]
                p=subprocess.run(cmd,capture_output=True,text=True,timeout=90)
                if p.returncode!=0: raise RuntimeError(p.stderr[-500:])
                result['preview']=[str(x) for x in sorted(preview_dir.glob('slide-*.jpg'))]
                result['pdf']=str(pdf)
            except (OSError,subprocess.TimeoutExpired,RuntimeError) as exc:
                warnings.append('预览失败: '+str(exc)[:200])
    return result
